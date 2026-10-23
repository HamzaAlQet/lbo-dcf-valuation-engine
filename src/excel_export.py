"""Formula-driven Excel export of the 3-statement, DCF and LBO model.

The workbook is a live model, not a dump of Python results. The only hard-coded
numbers are (a) data reported in the 10-K and (b) assumptions on the ``Inputs``
sheet. Every subtotal, projection, valuation and return is an Excel formula
(``=SUM(...)``, ``=MIN(...)``, ``=IRR(...)``), so a reviewer can trace and flex
every number.

Colour convention (banking standard)
------------------------------------
- **Blue** font: hard-coded input or reported historical figure.
- **Black** font: formula on the same sheet.
- **Green** font: formula linking to another sheet.

Design
------
Each sheet is first *declared* as a list of :class:`Line` objects whose cells are
callables. Row numbers are assigned on declaration; formulas are generated only at
render time, once every sheet's layout is known, so a formula may reference a row
on any sheet, including one declared later.

Every formula cell can also declare the value the Python engine computed for it.
:func:`write_workbook` returns these expectations so that the test-suite can
evaluate the workbook's formulas and prove that Excel and Python agree to the cent.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.modeling_engine import (
    HELD_FLAT_BALANCES,
    WORKING_CAPITAL_DRIVERS,
    ValuationModel,
    blended_cost_of_debt,
)

CellValue = Any
"""A number (input), a string starting with ``=`` (formula), text, or ``None``."""

# --------------------------------------------------------------------------------------
# Formatting
# --------------------------------------------------------------------------------------
NAVY, BLUE_HDR, SECTION_FILL, INPUT_FILL = "1F3864", "2F5597", "D9E1F2", "FFF2CC"
FONT_NAME = "Calibri"

NUM = '#,##0_);(#,##0);"-"_)'
NUM1 = '#,##0.0_);(#,##0.0);"-"_)'
PCT = '0.0%_);(0.0%);"-"_)'
PCT2 = '0.00%_);(0.00%);"-"_)'
MULT = '0.0"x"_);(0.0"x")'
PRICE = '$#,##0.00_);($#,##0.00)'
FACTOR = '0.0000_)'
DAYS = '0.0_)'
INT = '0_)'
TEXT = "@"

THIN = Side(style="thin", color="808080")
DOUBLE = Side(style="double", color="808080")

FIRST_YEAR_COL = 4   # column D
HEADER_ROW = 4       # numeric fiscal-year header (formatted as FY2015A / FY2016E)
FIRST_BODY_ROW = 6


def col_letter(index: int) -> str:
    """1-based column index -> letter."""
    return get_column_letter(index)


# --------------------------------------------------------------------------------------
# Declarative sheet model
# --------------------------------------------------------------------------------------
YearFn = Callable[[int], CellValue]
ScalarFn = Callable[[], CellValue]


@dataclass
class Line:
    """One row of a sheet.

    Attributes:
        key: Identifier used by formulas on any sheet (``None`` for headers/blanks).
        label: Text in column B.
        row: Worksheet row number.
        kind: ``"line"``, ``"section"`` or ``"blank"``.
        fmt: Default number format.
        style: ``"normal"``, ``"subtotal"``, ``"total"``, ``"memo"`` or ``"check"``.
        hist: Cell for each historical year column.
        proj: Cell for each projection year column.
        cells: Explicit cells by column letter: ``{col: (fn, fmt)}``.
        expect: Python-engine value for a column key (year or letter), if any.
        unit: Optional unit text for column C.
    """

    key: str | None
    label: str
    row: int
    kind: str = "line"
    fmt: str = NUM
    style: str = "normal"
    hist: YearFn | None = None
    proj: YearFn | None = None
    cells: dict[str, tuple[ScalarFn, str]] = field(default_factory=dict)
    expect: Callable[[int | str], float | None] | None = None
    unit: str = ""


class SheetSpec:
    """Layout of one worksheet, declared before any formula is rendered."""

    def __init__(self, book: WorkbookBuilder, name: str, title: str, *,
                 hist_years: list[int], proj_years: list[int],
                 hist_label: str = "Historical", proj_label: str = "Projected",
                 inputs_sheet: bool = False) -> None:
        self.book = book
        self.name = name
        self.title = title
        self.hist_years = hist_years
        self.proj_years = proj_years
        self.hist_label = hist_label
        self.proj_label = proj_label
        self.inputs_sheet = inputs_sheet
        self.lines: list[Line] = []
        self.rows: dict[str, int] = {}
        self.next_row = FIRST_BODY_ROW
        self.column_headers: dict[int, dict[str, str]] = {}  # row -> {col: text}

    # ---------------------------------------------------------------- declaring
    def section(self, label: str, columns: dict[str, str] | None = None) -> None:
        """Bold, shaded section header; ``columns`` adds header text in given columns."""
        if self.lines and self.lines[-1].kind != "blank":
            self.blank()
        self.lines.append(Line(None, label, self.next_row, kind="section"))
        if columns:
            self.column_headers[self.next_row] = columns
        self.next_row += 1

    def blank(self) -> None:
        """Empty spacer row."""
        self.lines.append(Line(None, "", self.next_row, kind="blank"))
        self.next_row += 1

    def line(self, key: str | None, label: str, *, hist: YearFn | None = None,
             proj: YearFn | None = None, value: ScalarFn | None = None,
             cells: dict[str, tuple[ScalarFn, str]] | None = None, fmt: str = NUM,
             style: str = "normal", unit: str = "",
             expect: Callable[[int | str], float | None] | None = None) -> None:
        """Declare a row. ``value`` is shorthand for a single cell in column D."""
        all_cells = dict(cells or {})
        if value is not None:
            all_cells["D"] = (value, fmt)
        if key is not None:
            if key in self.rows:
                raise ValueError(f"Duplicate key '{key}' on sheet {self.name}")
            self.rows[key] = self.next_row
        self.lines.append(Line(key, label, self.next_row, fmt=fmt, style=style, hist=hist,
                               proj=proj, cells=all_cells, expect=expect, unit=unit))
        self.next_row += 1

    # ---------------------------------------------------------------- referencing
    def r(self, key: str, year: int | None = None, *, col: str | None = None,
          sheet: str | None = None, absolute: bool = False) -> str:
        """Reference to a row on this or another sheet, from this sheet."""
        return self.book.ref(sheet or self.name, key, year, col=col, absolute=absolute,
                             from_sheet=self.name)

    def span(self, first: str, last: str, year: int | None = None, *,
             col: str | None = None) -> str:
        """Same-sheet range between two rows in one column, e.g. ``G10:G15``."""
        c = col or self.book.col(year)  # type: ignore[arg-type]
        return f"{c}{self.rows[first]}:{c}{self.rows[last]}"


class WorkbookBuilder:
    """Declares every sheet of the model, then renders them to an openpyxl workbook."""

    def __init__(self, model: ValuationModel) -> None:
        self.model = model
        self.schema = model.schema
        self.a = model.assumptions
        self.p = model.projection
        self.hist_years = self.schema.years
        self.bs_years = sorted(b.fiscal_year for b in self.schema.balance_sheets)
        self.proj_years = self.p.years
        self.base_year = self.p.base_year
        self.lbo_years = self.p.years[: self.a.transaction.holding_period_years]
        self.all_years = self.hist_years + self.proj_years
        self.sheets: dict[str, SheetSpec] = {}
        self.checks: list[tuple[str, str, str]] = []  # (description, sheet, cell)
        self.expected: dict[str, float] = {}

    # ---------------------------------------------------------------- helpers
    def col(self, year: int) -> str:
        """Column letter for a fiscal year (same on every time-series sheet)."""
        return col_letter(FIRST_YEAR_COL + self.all_years.index(year))

    def ref(self, sheet: str, key: str, year: int | None = None, *, col: str | None = None,
            absolute: bool = False, from_sheet: str | None = None) -> str:
        """A1 reference to ``key``'s row on ``sheet``, qualified when cross-sheet."""
        row = self.sheets[sheet].rows[key]
        c = col or self.col(year)  # type: ignore[arg-type]
        addr = f"${c}${row}" if absolute else f"{c}{row}"
        return addr if sheet == from_sheet else f"'{sheet}'!{addr}"

    def add(self, spec: SheetSpec) -> SheetSpec:
        """Register a sheet."""
        self.sheets[spec.name] = spec
        return spec

    # ---------------------------------------------------------------- build
    def build(self) -> Workbook:
        """Declare all sheets, then render them."""
        self._declare_inputs()
        self._declare_income_statement()
        self._declare_balance_sheet()
        self._declare_cash_flow()
        self._declare_debt()
        self._declare_dcf()
        self._declare_lbo()
        wb = Workbook()
        wb.remove(wb.active)
        cover = wb.create_sheet("Cover")
        for spec in self.sheets.values():
            self._render(wb.create_sheet(spec.name), spec)
        self._render_checks(wb.create_sheet("Checks"))
        self._render_cover(cover)
        return wb

    # ================================================================= Inputs
    def _declare_inputs(self) -> None:
        s = self.add(SheetSpec(self, "Inputs", "Assumptions & Inputs",
                               hist_years=self.hist_years, proj_years=self.proj_years,
                               hist_label="Historical (actual ratios)",
                               proj_label="Projection drivers", inputs_sheet=True))
        m, op, tx = self.a.market, self.a.operating, self.a.transaction
        sch = self.schema
        inc = sch.income_statement(self.base_year)

        s.section("Market & valuation inputs")
        s.line("share_price", "Current share price", value=lambda: m.share_price, fmt=PRICE)
        s.line("risk_free_rate", "Risk-free rate (10Y Treasury)",
               value=lambda: m.risk_free_rate, fmt=PCT2)
        s.line("equity_risk_premium", "Equity risk premium",
               value=lambda: m.equity_risk_premium, fmt=PCT2)
        s.line("levered_beta", "Levered beta", value=lambda: m.levered_beta, fmt=FACTOR)
        s.line("marginal_tax_rate", "Marginal tax rate (WACC tax shield)",
               value=lambda: m.marginal_tax_rate, fmt=PCT)
        s.line("terminal_growth_rate", "Terminal growth rate (Gordon)",
               value=lambda: m.terminal_growth_rate, fmt=PCT)
        s.line("dcf_exit_multiple", "Terminal EV / EBITDA multiple",
               value=lambda: m.exit_ev_ebitda_multiple, fmt=MULT)
        s.line("mid_year_convention", "Mid-year discounting (1 = on)",
               value=lambda: int(m.mid_year_convention), fmt=INT)
        s.line("deduct_repatriation_tax", "Deduct repatriation tax in equity bridge (1 = on)",
               value=lambda: int(m.deduct_repatriation_tax_in_equity_bridge), fmt=INT)
        s.line("repatriation_tax_rate", "Repatriation tax rate",
               value=lambda: (m.repatriation_tax_rate if m.repatriation_tax_rate is not None
                              else f"={s.r('unrecognized_repatriation_tax', col='D')}"
                                   f"/{s.r('undistributed_foreign_earnings', col='D')}"),
               fmt=PCT, expect=lambda _: sch.implied_repatriation_tax_rate
               if m.repatriation_tax_rate is None else None)
        s.line("days_in_year", "Days in year", value=lambda: 365, fmt=INT)

        s.section("Filing disclosures (10-K, latest fiscal year end)")
        s.line("shares_outstanding", "Shares outstanding at year end",
               value=lambda: sch.shares_outstanding or inc.diluted_shares, fmt=NUM1, unit="m")
        s.line("basic_shares_weighted", "Weighted-average basic shares",
               value=lambda: inc.basic_shares or inc.diluted_shares, fmt=NUM1, unit="m")
        s.line("diluted_shares_weighted", "Weighted-average diluted shares",
               value=lambda: inc.diluted_shares, fmt=NUM1, unit="m")
        s.line("diluted_shares_outstanding", "Diluted shares (period-end + dilution)",
               value=lambda: f"={s.r('shares_outstanding', col='D')}"
                             f"+{s.r('diluted_shares_weighted', col='D')}"
                             f"-{s.r('basic_shares_weighted', col='D')}",
               fmt=NUM1, unit="m", expect=lambda _: sch.diluted_shares_outstanding)
        s.line("foreign_cash", "Cash & securities held by foreign subsidiaries",
               value=lambda: sch.foreign_cash_and_investments or 0.0)
        s.line("undistributed_foreign_earnings", "Undistributed foreign earnings",
               value=lambda: sch.undistributed_foreign_earnings or 0.0)
        s.line("unrecognized_repatriation_tax", "Unrecognized deferred tax on foreign earnings",
               value=lambda: sch.unrecognized_repatriation_tax or 0.0)
        s.line("portfolio_yield", "Yield on cash & marketable securities",
               value=lambda: sch.investment_portfolio_yield or 0.0, fmt=PCT2)

        s.section("Operating drivers")
        IS, BS, CF = "IS", "BS", "CFS"

        def ratio(num: str, den: str, num_sheet: str = IS, den_sheet: str = IS,
                  sign: str = "") -> YearFn:
            return lambda y: (f"={sign}{s.r(num, y, sheet=num_sheet)}"
                              f"/{s.r(den, y, sheet=den_sheet)}")

        def driver(key: str) -> YearFn:
            return lambda y: float(self.p.drivers.at[key, y])

        first_hist = self.hist_years[0]
        s.line("revenue_growth", "Revenue growth", fmt=PCT,
               hist=lambda y: None if y == first_hist else
               f"={s.r('net_sales', y, sheet=IS)}/{s.r('net_sales', y - 1, sheet=IS)}-1",
               proj=driver("revenue_growth"))
        s.line("gross_margin", "Gross margin", fmt=PCT,
               hist=ratio("gross_margin", "net_sales"), proj=driver("gross_margin"))
        s.line("rd_pct_sales", "R&D, % of sales", fmt=PCT,
               hist=ratio("research_and_development", "net_sales"), proj=driver("rd_pct_sales"))
        s.line("sga_pct_sales", "SG&A, % of sales", fmt=PCT,
               hist=ratio("selling_general_administrative", "net_sales"),
               proj=driver("sga_pct_sales"))
        s.line("da_pct_sales", "D&A, % of sales", fmt=PCT,
               hist=ratio("depreciation_and_amortization", "net_sales"),
               proj=driver("da_pct_sales"))
        s.line("capex_pct_sales", "Capex, % of sales", fmt=PCT,
               hist=ratio("capital_expenditures", "net_sales", num_sheet=CF, sign="-"),
               proj=driver("capex_pct_sales"))
        s.line("tax_rate", "Tax rate", fmt=PCT,
               hist=ratio("income_tax_expense", "pretax_income"), proj=driver("tax_rate"))

        base_col = self.col(self.base_year)
        for key, (bs_field, base, kind) in WORKING_CAPITAL_DRIVERS.items():
            label = key.replace("_", " ").replace("pct", "%").capitalize()
            override = getattr(op, key)
            days = f"*{s.r('days_in_year', col='D', absolute=True)}" if kind == "days" else ""

            def hist_fn(y: int, f: str = bs_field, b: str = base, d: str = days) -> CellValue:
                if y not in self.bs_years:
                    return None
                return f"={s.r(f, y, sheet=BS)}/{s.r(b, y, sheet=IS)}{d}"

            def proj_fn(y: int, k: str = key, o: float | None = override) -> CellValue:
                return o if o is not None else f"={s.r(k, col=base_col, absolute=True)}"

            s.line(key, label, fmt=DAYS if kind == "days" else PCT, hist=hist_fn,
                   proj=proj_fn, expect=lambda y, k=key: (
                       float(self.p.drivers.at[k, y]) if y in self.proj_years else None))

        s.line("interest_income_yield", "Interest income yield on cash & securities", fmt=PCT2,
               proj=lambda y: (op.interest_income_yield if op.interest_income_yield is not None
                               else f"={s.r('portfolio_yield', col='D', absolute=True)}"),
               expect=lambda y: float(self.p.drivers.at["interest_income_yield", y]))
        s.line("cost_of_debt", "Pre-tax cost of debt", fmt=PCT2,
               proj=lambda y: (m.pre_tax_cost_of_debt if m.pre_tax_cost_of_debt is not None
                               else f"={s.r('blended_cost_of_debt', col='D', sheet='Debt', absolute=True)}"),
               expect=lambda y: blended_cost_of_debt(sch, m))
        s.line("dividend_payout_ratio", "Dividends, % of net income", fmt=PCT,
               hist=ratio("dividends_paid", "net_income", num_sheet=CF, sign="-"),
               proj=lambda y: op.dividend_payout_ratio)
        s.line("buyback_pct_net_income", "Share repurchases, % of net income", fmt=PCT,
               hist=ratio("share_repurchases", "net_income", num_sheet=CF, sign="-"),
               proj=lambda y: op.buyback_pct_net_income)

        s.section("Transaction assumptions (LBO)")
        s.line("entry_multiple", "Entry EV / LTM EBITDA",
               value=lambda: tx.entry_ev_ebitda_multiple, fmt=MULT)
        s.line("exit_multiple", "Exit EV / EBITDA", value=lambda: tx.exit_ev_ebitda_multiple,
               fmt=MULT)
        s.line("holding_period", "Holding period (years)",
               value=lambda: tx.holding_period_years, fmt=INT)
        s.line("debt_pct_of_entry_ev", "Total debt, % of entry EV",
               value=lambda: tx.debt_pct_of_entry_ev, fmt=PCT)
        s.line("cash_sweep_pct", "Excess cash flow sweep", value=lambda: tx.cash_sweep_pct,
               fmt=PCT)
        s.line("minimum_cash", "Minimum operating cash", value=lambda: tx.minimum_cash)
        s.line("advisory_fee_pct", "Advisory & transaction fees, % of EV",
               value=lambda: tx.advisory_fee_pct_ev, fmt=PCT2)
        s.line("financing_fee_pct", "Financing fees, % of debt",
               value=lambda: tx.financing_fee_pct_debt, fmt=PCT2)
        s.line("foreign_cash_usable_pct", "Offshore cash repatriated at close",
               value=lambda: tx.foreign_cash_usable_pct, fmt=PCT)
        s.line("max_leverage_covenant", "Covenant: maximum Total debt / EBITDA",
               value=lambda: tx.max_leverage_covenant, fmt=MULT)
        s.line("min_coverage_covenant", "Covenant: minimum EBITDA / Interest",
               value=lambda: tx.min_interest_coverage_covenant, fmt=MULT)
        for i, t in enumerate(tx.tranches):
            s.section(f"Tranche {i + 1}: {t.name} ({t.seniority.value.replace('_', ' ')})")
            s.line(f"tranche{i}_share", "Share of total debt", value=lambda t=t: t.share_of_debt,
                   fmt=PCT)
            s.line(f"tranche{i}_rate", "Cash interest rate", value=lambda t=t: t.interest_rate,
                   fmt=PCT2)
            s.line(f"tranche{i}_amort", "Mandatory amortization, % of original",
                   value=lambda t=t: t.mandatory_amortization_pct, fmt=PCT)
            s.line(f"tranche{i}_prepayable", "Prepayable from sweep (1 = yes)",
                   value=lambda t=t: int(t.prepayable), fmt=INT)

    # ================================================================= Income statement
    def _declare_income_statement(self) -> None:
        s = self.add(SheetSpec(self, "IS", "Income Statement",
                               hist_years=self.hist_years, proj_years=self.proj_years))
        sch, inc = self.schema, self.schema.income_statement
        inp = lambda key, y: s.r(key, y, sheet="Inputs")  # noqa: E731
        source = lambda key: (lambda y: self._expect_income(key, y))  # noqa: E731

        def reported(field: str) -> YearFn:
            return lambda y: getattr(inc(y), field)

        s.section("Income statement")
        s.line("net_sales", "Net sales", style="subtotal", hist=reported("net_sales"),
               proj=lambda y: f"={s.r('net_sales', y - 1)}*(1+{inp('revenue_growth', y)})",
               expect=source("net_sales"))
        s.line("cost_of_sales", "Cost of sales", hist=reported("cost_of_sales"),
               proj=lambda y: f"={s.r('net_sales', y)}*(1-{inp('gross_margin', y)})",
               expect=source("cost_of_sales"))
        both = lambda fn: {"hist": fn, "proj": fn}  # noqa: E731
        s.line("gross_margin", "Gross profit", style="subtotal",
               **both(lambda y: f"={s.r('net_sales', y)}-{s.r('cost_of_sales', y)}"),
               expect=source("gross_margin"))
        s.line("gross_margin_pct", "  Gross margin %", fmt=PCT, style="memo",
               **both(lambda y: f"={s.r('gross_margin', y)}/{s.r('net_sales', y)}"))
        s.line("research_and_development", "Research and development",
               hist=reported("research_and_development"),
               proj=lambda y: f"={s.r('net_sales', y)}*{inp('rd_pct_sales', y)}",
               expect=source("research_and_development"))
        s.line("selling_general_administrative", "Selling, general and administrative",
               hist=reported("selling_general_administrative"),
               proj=lambda y: f"={s.r('net_sales', y)}*{inp('sga_pct_sales', y)}",
               expect=source("selling_general_administrative"))
        s.line("total_operating_expenses", "Total operating expenses", style="subtotal",
               **both(lambda y: f"=SUM({s.span('research_and_development', 'selling_general_administrative', y)})"),
               expect=source("total_operating_expenses"))
        s.line("operating_income", "Operating income (EBIT)", style="subtotal",
               **both(lambda y: f"={s.r('gross_margin', y)}-{s.r('total_operating_expenses', y)}"),
               expect=source("operating_income"))
        s.line("operating_margin", "  EBIT margin %", fmt=PCT, style="memo",
               **both(lambda y: f"={s.r('operating_income', y)}/{s.r('net_sales', y)}"))
        s.line("depreciation_and_amortization", "Depreciation & amortization", style="memo",
               hist=lambda y: f"={s.r('depreciation_and_amortization', y, sheet='CFS')}",
               proj=lambda y: f"={s.r('net_sales', y)}*{inp('da_pct_sales', y)}",
               expect=source("depreciation_and_amortization"))
        s.line("ebitda", "EBITDA", style="subtotal",
               **both(lambda y: f"={s.r('operating_income', y)}+{s.r('depreciation_and_amortization', y)}"),
               expect=lambda y: sch.ebitda(y) if y in self.hist_years
               else self.p.line("income", "ebitda", y))
        s.line("ebitda_margin", "  EBITDA margin %", fmt=PCT, style="memo",
               **both(lambda y: f"={s.r('ebitda', y)}/{s.r('net_sales', y)}"))

        def opening(field: str, y: int) -> str:
            return s.r(field, y - 1, sheet="BS")

        s.line("interest_income", "Interest income",
               proj=lambda y: f"={inp('interest_income_yield', y)}*({opening('cash_and_equivalents', y)}"
                              f"+{opening('short_term_investments', y)}"
                              f"+{opening('long_term_investments', y)})",
               expect=source("interest_income"))
        s.line("interest_expense", "Interest expense",
               proj=lambda y: f"={inp('cost_of_debt', y)}*({opening('commercial_paper', y)}"
                              f"+{opening('current_portion_long_term_debt', y)}"
                              f"+{opening('long_term_debt', y)})",
               expect=source("interest_expense"))
        s.line("other_income_expense", "Other income / (expense), net",
               hist=reported("other_income_expense"),
               proj=lambda y: f"={s.r('interest_income', y)}-{s.r('interest_expense', y)}",
               expect=source("other_income_expense"))
        s.line("pretax_income", "Income before taxes", style="subtotal",
               **both(lambda y: f"={s.r('operating_income', y)}+{s.r('other_income_expense', y)}"),
               expect=source("pretax_income"))
        s.line("income_tax_expense", "Provision for income taxes",
               hist=reported("income_tax_expense"),
               proj=lambda y: f"={s.r('pretax_income', y)}*{inp('tax_rate', y)}",
               expect=source("income_tax_expense"))
        s.line("net_income", "Net income", style="total",
               **both(lambda y: f"={s.r('pretax_income', y)}-{s.r('income_tax_expense', y)}"),
               expect=source("net_income"))
        s.blank()
        s.line("diluted_shares", "Diluted shares", fmt=NUM1, unit="m",
               hist=reported("diluted_shares"),
               proj=lambda y: f"={s.r('diluted_shares_outstanding', col='D', sheet='Inputs', absolute=True)}",
               expect=source("diluted_shares"))
        s.line("eps_diluted", "Diluted EPS", fmt=PRICE,
               **both(lambda y: f"={s.r('net_income', y)}/{s.r('diluted_shares', y)}"),
               expect=lambda y: self.p.line("income", "eps_diluted", y)
               if y in self.proj_years else None)

        # Reported segment data, with a tie-out to consolidated net sales.
        for seg_type in ("geographic", "product"):
            segments = [x for x in sch.revenue_segments if x.segment_type == seg_type]
            if not segments:
                continue
            s.section(f"Net sales by {'segment' if seg_type == 'geographic' else 'product'} (reported)")
            keys = []
            for seg in segments:
                key = f"seg_{seg_type}_{len(keys)}"
                keys.append(key)
                s.line(key, seg.name, hist=lambda y, seg=seg: seg.net_sales.get(y))
            s.line(f"seg_{seg_type}_total", "Total", style="subtotal",
                   hist=lambda y, k=keys: f"=SUM({s.span(k[0], k[-1], y)})")
            s.line(f"seg_{seg_type}_check", "Check: segments - net sales", style="check",
                   hist=lambda y, t=seg_type: f"={s.r(f'seg_{t}_total', y)}-{s.r('net_sales', y)}")

    def _expect_income(self, key: str, year: int) -> float | None:
        """Python value for an income-statement row (reported or projected)."""
        if year in self.proj_years:
            return self.p.line("income", key, year) if key in self.p.income.index else None
        if key == "depreciation_and_amortization":
            return self.schema.cash_flow_statement(year).depreciation_and_amortization
        return getattr(self.schema.income_statement(year), key, None)

    # ================================================================= Balance sheet
    def _declare_balance_sheet(self) -> None:
        s = self.add(SheetSpec(self, "BS", "Balance Sheet",
                               hist_years=self.bs_years, proj_years=self.proj_years))
        sch = self.schema
        inp = lambda key, y: s.r(key, y, sheet="Inputs")  # noqa: E731
        is_ = lambda key, y: s.r(key, y, sheet="IS")  # noqa: E731
        cfs = lambda key, y: s.r(key, y, sheet="CFS")  # noqa: E731
        days = s.r("days_in_year", col="D", sheet="Inputs", absolute=True)
        wc = {f: (k, b, kind) for k, (f, b, kind) in WORKING_CAPITAL_DRIVERS.items()}

        def expect(key: str) -> Callable[[int | str], float | None]:
            def fn(y: int | str) -> float | None:
                if y in self.proj_years:
                    return self.p.line("balance", key, y)  # type: ignore[arg-type]
                return getattr(sch.balance_sheet(y), key, None)  # type: ignore[arg-type]
            return fn

        def item(key: str, label: str, style: str = "normal",
                 proj: YearFn | None = None, hist: YearFn | None = None) -> None:
            if proj is None and key in HELD_FLAT_BALANCES:
                proj = lambda y: f"={s.r(key, y - 1)}"  # noqa: E731
            if proj is None and key in wc:
                drv, base, kind = wc[key]
                scale = f"/{days}" if kind == "days" else ""
                proj = lambda y: f"={is_(base, y)}*{inp(drv, y)}{scale}"  # noqa: E731
            if hist is None:
                hist = lambda y: getattr(sch.balance_sheet(y), key)  # noqa: E731
            s.line(key, label, style=style, hist=hist, proj=proj, expect=expect(key))

        both = lambda fn: {"hist": fn, "proj": fn}  # noqa: E731

        s.section("Assets")
        item("cash_and_equivalents", "Cash and cash equivalents",
             proj=lambda y: f"={cfs('ending_cash', y)}")
        item("short_term_investments", "Short-term marketable securities")
        item("accounts_receivable", "Accounts receivable")
        item("inventories", "Inventories")
        item("deferred_tax_assets_current", "Deferred tax assets")
        item("vendor_non_trade_receivables", "Vendor non-trade receivables")
        item("other_current_assets", "Other current assets")
        item("total_current_assets", "Total current assets", style="subtotal",
             **both(lambda y: f"=SUM({s.span('cash_and_equivalents', 'other_current_assets', y)})"))
        item("long_term_investments", "Long-term marketable securities")
        item("ppe_net", "Property, plant and equipment, net",
             proj=lambda y: f"={s.r('ppe_net', y - 1)}-{cfs('capital_expenditures', y)}"
                            f"-{is_('depreciation_and_amortization', y)}")
        item("goodwill", "Goodwill")
        item("intangible_assets", "Acquired intangible assets, net")
        item("other_non_current_assets", "Other assets")
        item("total_assets", "Total assets", style="total",
             **both(lambda y: f"={s.r('total_current_assets', y)}"
                              f"+SUM({s.span('long_term_investments', 'other_non_current_assets', y)})"))

        s.section("Liabilities")
        item("accounts_payable", "Accounts payable")
        item("accrued_expenses", "Accrued expenses")
        item("deferred_revenue_current", "Deferred revenue")
        item("commercial_paper", "Commercial paper")
        item("current_portion_long_term_debt", "Current portion of long-term debt")
        item("total_current_liabilities", "Total current liabilities", style="subtotal",
             **both(lambda y: f"=SUM({s.span('accounts_payable', 'current_portion_long_term_debt', y)})"))
        item("deferred_revenue_non_current", "Deferred revenue, non-current")
        item("long_term_debt", "Long-term debt")
        item("other_non_current_liabilities", "Other non-current liabilities")
        item("total_liabilities", "Total liabilities", style="subtotal",
             **both(lambda y: f"={s.r('total_current_liabilities', y)}"
                              f"+SUM({s.span('deferred_revenue_non_current', 'other_non_current_liabilities', y)})"))

        s.section("Shareholders' equity")
        item("common_stock_and_apic", "Common stock and additional paid-in capital")
        item("retained_earnings", "Retained earnings",
             proj=lambda y: f"={s.r('retained_earnings', y - 1)}+{is_('net_income', y)}"
                            f"+{cfs('dividends_paid', y)}+{cfs('share_repurchases', y)}")
        item("accumulated_other_comprehensive_income", "Accumulated other comprehensive income")
        item("total_equity", "Total shareholders' equity", style="subtotal",
             **both(lambda y: f"=SUM({s.span('common_stock_and_apic', 'accumulated_other_comprehensive_income', y)})"))
        item("total_liabilities_and_equity", "Total liabilities and shareholders' equity",
             style="total",
             **both(lambda y: f"={s.r('total_liabilities', y)}+{s.r('total_equity', y)}"))
        s.line("balance_check", "Check: Assets - (Liabilities + Equity)", style="check",
               **both(lambda y: f"={s.r('total_assets', y)}-{s.r('total_liabilities_and_equity', y)}"))

        s.section("Memo")
        s.line("net_working_capital", "Net working capital (operating)",
               **both(lambda y: f"=({s.r('total_current_assets', y)}-{s.r('cash_and_equivalents', y)}"
                                f"-{s.r('short_term_investments', y)}-{s.r('deferred_tax_assets_current', y)})"
                                f"-({s.r('total_current_liabilities', y)}-{s.r('commercial_paper', y)}"
                                f"-{s.r('current_portion_long_term_debt', y)})"),
               expect=lambda y: (self.p.statements[y][1] if y in self.proj_years
                                 else sch.balance_sheet(y)).net_working_capital)  # type: ignore[index,arg-type]
        s.line("total_debt", "Total debt",
               **both(lambda y: f"={s.r('commercial_paper', y)}+{s.r('current_portion_long_term_debt', y)}"
                                f"+{s.r('long_term_debt', y)}"))
        s.line("cash_and_investments", "Cash and marketable securities",
               **both(lambda y: f"={s.r('cash_and_equivalents', y)}+{s.r('short_term_investments', y)}"
                                f"+{s.r('long_term_investments', y)}"))

    # ================================================================= Cash flow
    def _declare_cash_flow(self) -> None:
        s = self.add(SheetSpec(self, "CFS", "Cash Flow Statement",
                               hist_years=self.hist_years, proj_years=self.proj_years))
        sch = self.schema
        is_ = lambda key, y: s.r(key, y, sheet="IS")  # noqa: E731
        bs = lambda key, y: s.r(key, y, sheet="BS")  # noqa: E731
        inp = lambda key, y: s.r(key, y, sheet="Inputs")  # noqa: E731
        both = lambda fn: {"hist": fn, "proj": fn}  # noqa: E731

        def expect(key: str) -> Callable[[int | str], float | None]:
            def fn(y: int | str) -> float | None:
                if y in self.proj_years:
                    return float(self.p.cash_flow.at[key, y])
                return getattr(sch.cash_flow_statement(y), key, None)  # type: ignore[arg-type]
            return fn

        def item(key: str, label: str, proj: YearFn | None = None, style: str = "normal",
                 hist: YearFn | None = None) -> None:
            if hist is None:
                hist = lambda y: getattr(sch.cash_flow_statement(y), key)  # noqa: E731
            s.line(key, label, hist=hist, proj=proj, style=style, expect=expect(key))

        def asset_change(field: str) -> YearFn:
            return lambda y: f"={bs(field, y - 1)}-{bs(field, y)}"

        def liability_change(field: str) -> YearFn:
            return lambda y: f"={bs(field, y)}-{bs(field, y - 1)}"

        item("beginning_cash", "Cash, beginning of year",
             proj=lambda y: f"={bs('cash_and_equivalents', y - 1)}")
        s.section("Operating activities")
        item("net_income", "Net income", proj=lambda y: f"={is_('net_income', y)}")
        item("depreciation_and_amortization", "Depreciation and amortization",
             proj=lambda y: f"={is_('depreciation_and_amortization', y)}")
        item("share_based_compensation", "Share-based compensation")
        item("deferred_income_taxes", "Deferred income taxes")
        item("other_non_cash_items", "Other non-cash items")
        item("change_accounts_receivable", "(Increase) / decrease in receivables",
             proj=asset_change("accounts_receivable"))
        item("change_inventories", "(Increase) / decrease in inventories",
             proj=asset_change("inventories"))
        item("change_vendor_non_trade_receivables", "(Increase) / decrease in vendor receivables",
             proj=asset_change("vendor_non_trade_receivables"))
        item("change_other_assets", "(Increase) / decrease in other assets",
             proj=asset_change("other_current_assets"))
        item("change_accounts_payable", "Increase / (decrease) in payables",
             proj=liability_change("accounts_payable"))
        item("change_deferred_revenue", "Increase / (decrease) in deferred revenue",
             proj=liability_change("deferred_revenue_current"))
        item("change_other_liabilities", "Increase / (decrease) in other liabilities",
             proj=liability_change("accrued_expenses"))
        item("cash_from_operations", "Cash from operating activities", style="subtotal",
             **both(lambda y: f"=SUM({s.span('net_income', 'change_other_liabilities', y)})"))

        s.section("Investing activities")
        item("purchases_of_investments", "Purchases of marketable securities")
        item("maturities_of_investments", "Maturities of marketable securities")
        item("sales_of_investments", "Sales of marketable securities")
        item("acquisitions", "Business acquisitions, net")
        item("capital_expenditures", "Capital expenditures",
             proj=lambda y: f"=-{is_('net_sales', y)}*{inp('capex_pct_sales', y)}")
        item("purchases_of_intangibles", "Purchases of intangible assets")
        item("other_investing", "Other investing")
        item("cash_from_investing", "Cash from investing activities", style="subtotal",
             **both(lambda y: f"=SUM({s.span('purchases_of_investments', 'other_investing', y)})"))

        s.section("Financing activities")
        item("proceeds_from_stock_issuance", "Proceeds from issuance of common stock")
        item("excess_tax_benefits", "Excess tax benefits from equity awards")
        item("taxes_on_net_share_settlement", "Taxes on net share settlement")
        item("dividends_paid", "Dividends paid",
             proj=lambda y: f"=-{is_('net_income', y)}*{inp('dividend_payout_ratio', y)}")
        item("share_repurchases", "Repurchases of common stock",
             proj=lambda y: f"=-{is_('net_income', y)}*{inp('buyback_pct_net_income', y)}")
        item("proceeds_from_debt_net", "Proceeds from issuance of debt, net")
        item("repayment_of_debt", "Repayment of debt")
        item("change_in_commercial_paper", "Change in commercial paper, net")
        item("other_financing", "Other financing")
        item("cash_from_financing", "Cash from financing activities", style="subtotal",
             **both(lambda y: f"=SUM({s.span('proceeds_from_stock_issuance', 'other_financing', y)})"))

        s.blank()
        item("net_change_in_cash", "Net change in cash", style="subtotal",
             **both(lambda y: f"={s.r('cash_from_operations', y)}+{s.r('cash_from_investing', y)}"
                              f"+{s.r('cash_from_financing', y)}"))
        item("ending_cash", "Cash, end of year", style="total",
             **both(lambda y: f"={s.r('beginning_cash', y)}+{s.r('net_change_in_cash', y)}"))
        s.line("cash_check", "Check: ending cash - balance sheet cash", style="check",
               hist=lambda y: (f"={s.r('ending_cash', y)}-{bs('cash_and_equivalents', y)}"
                               if y in self.bs_years else None),
               proj=lambda y: f"={s.r('ending_cash', y)}-{bs('cash_and_equivalents', y)}")
        s.line("net_income_check", "Check: CFS net income - IS net income", style="check",
               hist=lambda y: f"={s.r('net_income', y)}-{is_('net_income', y)}")

    # ================================================================= Debt
    def _declare_debt(self) -> None:
        s = self.add(SheetSpec(self, "Debt", "Existing Debt Schedule (10-K)",
                               hist_years=[], proj_years=[]))
        ds = self.schema.debt_schedule
        s.section("Term debt tranches", columns={
            "C": "Currency", "D": "Principal", "E": "Eff. rate low", "F": "Eff. rate high",
            "G": "Eff. rate mid", "H": "Maturity from", "I": "Maturity to"})
        tranches = ds.tranches if ds else []
        for i, t in enumerate(tranches):
            mid = lambda i=i: f"=({s.r(f'tr{i}', col='E')}+{s.r(f'tr{i}', col='F')})/2"
            s.line(f"tr{i}", t.name, cells={
                "C": (lambda t=t: t.currency, TEXT),
                "D": (lambda t=t: t.principal, NUM),
                "E": (lambda t=t: t.effective_rate_low, PCT2),
                "F": (lambda t=t: t.effective_rate_high, PCT2),
                "G": (mid, PCT2),
                "H": (lambda t=t: t.maturity_year_start, INT),
                "I": (lambda t=t: t.maturity_year_end, INT),
            }, expect=lambda c, t=t: t.effective_rate if c == "G" else None)
        first, last = (f"tr0", f"tr{len(tranches) - 1}") if tranches else (None, None)
        rng = (lambda c: f"{c}{s.rows[first]}:{c}{s.rows[last]}") if tranches else None

        s.section("Reconciliation to balance sheet")
        s.line("total_principal", "Sum of tranche principal", style="subtotal",
               value=lambda: f"=SUM({rng('D')})" if rng else 0.0,
               expect=lambda _: ds.total_principal if ds else 0.0)
        s.line("reported_total_term_debt", "Total term debt (reported)",
               value=lambda: ds.reported_total_term_debt if ds else 0.0)
        s.line("principal_check", "Check: tranches - reported total", style="check",
               value=lambda: f"={s.r('total_principal', col='D')}-{s.r('reported_total_term_debt', col='D')}")
        s.line("unamortized_discount", "Unamortized discount",
               value=lambda: ds.unamortized_discount if ds else 0.0)
        s.line("hedge_adjustments", "Hedge accounting fair value adjustments",
               value=lambda: ds.hedge_accounting_adjustments if ds else 0.0)
        s.line("carrying_value", "Carrying value of term debt", style="subtotal",
               value=lambda: f"={s.r('reported_total_term_debt', col='D')}"
                             f"+{s.r('unamortized_discount', col='D')}"
                             f"+{s.r('hedge_adjustments', col='D')}",
               expect=lambda _: ds.carrying_value if ds else 0.0)
        bs_col = self.col(self.base_year)
        s.line("bs_term_debt", "Balance sheet: current + long-term debt",
               value=lambda: f"={s.r('current_portion_long_term_debt', col=bs_col, sheet='BS')}"
                             f"+{s.r('long_term_debt', col=bs_col, sheet='BS')}")
        s.line("carrying_check", "Check: carrying value - balance sheet", style="check",
               value=lambda: f"={s.r('carrying_value', col='D')}-{s.r('bs_term_debt', col='D')}")
        s.line("commercial_paper", "Commercial paper outstanding",
               value=lambda: f"={s.r('commercial_paper', col=bs_col, sheet='BS')}")
        s.line("commercial_paper_rate", "Commercial paper weighted-average rate",
               value=lambda: (ds.commercial_paper_rate or 0.0) if ds else 0.0, fmt=PCT2)
        s.line("term_debt_rate", "Term debt weighted-average effective rate", fmt=PCT2,
               value=lambda: (f"=SUMPRODUCT({rng('D')},{rng('G')})/{s.r('total_principal', col='D')}"
                              if rng else 0.0),
               expect=lambda _: ds.weighted_average_effective_rate if ds else None)
        s.line("blended_cost_of_debt", "Blended pre-tax cost of debt", style="total", fmt=PCT2,
               value=lambda: (f"=(SUMPRODUCT({rng('D')},{rng('G')})"
                              f"+{s.r('commercial_paper', col='D')}*{s.r('commercial_paper_rate', col='D')})"
                              f"/({s.r('total_principal', col='D')}+{s.r('commercial_paper', col='D')})"
                              if rng else 0.0),
               expect=lambda _: blended_cost_of_debt(self.schema, self.a.market))

    # ================================================================= DCF
    def _declare_dcf(self) -> None:
        s = self.add(SheetSpec(self, "DCF", "Discounted Cash Flow Valuation",
                               hist_years=[self.base_year], proj_years=self.proj_years,
                               hist_label="Actual"))
        d = self.model.dcf
        w = d.wacc
        base = self.col(self.base_year)
        inp = lambda key: s.r(key, col="D", sheet="Inputs", absolute=True)  # noqa: E731
        D = lambda key: s.r(key, col="D", absolute=True)  # noqa: E731
        first_proj = self.col(self.proj_years[0])
        last = self.proj_years[-1]

        s.section("Weighted average cost of capital")
        s.line("share_price", "Share price", value=lambda: f"={inp('share_price')}", fmt=PRICE)
        s.line("diluted_shares", "Diluted shares (m)", fmt=NUM1,
               value=lambda: f"={inp('diluted_shares_outstanding')}")
        s.line("market_value_of_equity", "Market value of equity",
               value=lambda: f"={D('share_price')}*{D('diluted_shares')}",
               expect=lambda _: w.market_value_of_equity)
        s.line("total_debt", "Total debt (book)",
               value=lambda: f"={s.r('total_debt', col=base, sheet='BS')}",
               expect=lambda _: w.total_debt)
        s.line("total_capitalization", "Total capitalization", style="subtotal",
               value=lambda: f"={D('market_value_of_equity')}+{D('total_debt')}")
        s.line("equity_weight", "Equity / total capitalization", fmt=PCT,
               value=lambda: f"={D('market_value_of_equity')}/{D('total_capitalization')}",
               expect=lambda _: w.equity_weight)
        s.line("debt_weight", "Debt / total capitalization", fmt=PCT,
               value=lambda: f"={D('total_debt')}/{D('total_capitalization')}",
               expect=lambda _: w.debt_weight)
        s.line("risk_free_rate", "Risk-free rate", fmt=PCT2,
               value=lambda: f"={inp('risk_free_rate')}")
        s.line("levered_beta", "Levered beta", fmt=FACTOR, value=lambda: f"={inp('levered_beta')}")
        s.line("equity_risk_premium", "Equity risk premium", fmt=PCT2,
               value=lambda: f"={inp('equity_risk_premium')}")
        s.line("cost_of_equity", "Cost of equity (CAPM)", fmt=PCT2, style="subtotal",
               value=lambda: f"={D('risk_free_rate')}+{D('levered_beta')}*{D('equity_risk_premium')}",
               expect=lambda _: w.cost_of_equity)
        s.line("pre_tax_cost_of_debt", "Pre-tax cost of debt", fmt=PCT2,
               value=lambda: f"={s.r('cost_of_debt', col=first_proj, sheet='Inputs', absolute=True)}",
               expect=lambda _: w.pre_tax_cost_of_debt)
        s.line("tax_rate", "Marginal tax rate", fmt=PCT,
               value=lambda: f"={inp('marginal_tax_rate')}")
        s.line("after_tax_cost_of_debt", "After-tax cost of debt", fmt=PCT2,
               value=lambda: f"={D('pre_tax_cost_of_debt')}*(1-{D('tax_rate')})",
               expect=lambda _: w.after_tax_cost_of_debt)
        s.line("wacc", "WACC", fmt=PCT2, style="total",
               value=lambda: f"={D('equity_weight')}*{D('cost_of_equity')}"
                             f"+{D('debt_weight')}*{D('after_tax_cost_of_debt')}",
               expect=lambda _: w.wacc)

        sched = d.schedule
        exp = lambda key: (lambda y: float(sched.at[key, y]))  # noqa: E731
        s.section("Unlevered free cash flow")
        s.line("operating_income", "EBIT",
               proj=lambda y: f"={s.r('operating_income', y, sheet='IS')}",
               expect=exp("operating_income"))
        s.line("tax_rate_proj", "Tax rate", fmt=PCT,
               proj=lambda y: f"={s.r('tax_rate', y, sheet='Inputs')}")
        s.line("nopat", "NOPAT = EBIT x (1 - t)",
               proj=lambda y: f"={s.r('operating_income', y)}*(1-{s.r('tax_rate_proj', y)})",
               expect=exp("nopat"))
        s.line("depreciation_and_amortization", "Plus: D&A",
               proj=lambda y: f"={s.r('depreciation_and_amortization', y, sheet='IS')}",
               expect=exp("depreciation_and_amortization"))
        s.line("capex", "Less: capital expenditures",
               proj=lambda y: f"=-{s.r('capital_expenditures', y, sheet='CFS')}",
               expect=exp("capex"))
        s.line("change_in_nwc", "Less: increase in net working capital",
               proj=lambda y: f"={s.r('net_working_capital', y, sheet='BS')}"
                              f"-{s.r('net_working_capital', y - 1, sheet='BS')}",
               expect=exp("change_in_nwc"))
        s.line("unlevered_fcf", "Unlevered free cash flow", style="total",
               proj=lambda y: f"={s.r('nopat', y)}+{s.r('depreciation_and_amortization', y)}"
                              f"-{s.r('capex', y)}-{s.r('change_in_nwc', y)}",
               expect=exp("unlevered_fcf"))
        s.line("period_index", "Year", fmt=INT,
               proj=lambda y: f"={self.col(y)}${HEADER_ROW}-${base}${HEADER_ROW}")
        s.line("discount_period", "Discount period", fmt=DAYS,
               proj=lambda y: f"=IF({inp('mid_year_convention')}=1,{s.r('period_index', y)}-0.5,"
                              f"{s.r('period_index', y)})",
               expect=exp("discount_period"))
        s.line("discount_factor", "Discount factor", fmt=FACTOR,
               proj=lambda y: f"=1/(1+{D('wacc')})^{s.r('discount_period', y)}",
               expect=exp("discount_factor"))
        s.line("pv_fcf", "PV of unlevered FCF", style="subtotal",
               proj=lambda y: f"={s.r('unlevered_fcf', y)}*{s.r('discount_factor', y)}",
               expect=exp("pv_fcf"))

        g, x = d.gordon, d.exit
        G = lambda key: s.r(key, col="D")  # noqa: E731
        E = lambda key: s.r(key, col="E")  # noqa: E731
        pv_range = f"{first_proj}{{row}}:{self.col(last)}{{row}}"

        def pair(key: str, label: str, gordon: ScalarFn, exit_: ScalarFn, fmt: str = NUM,
                 attr: str | None = None, style: str = "normal",
                 fmts: tuple[str, str] | None = None) -> None:
            f1, f2 = fmts or (fmt, fmt)
            s.line(key, label, cells={"D": (gordon, f1), "E": (exit_, f2)}, fmt=fmt, style=style,
                   expect=(lambda c: getattr(g if c == "D" else x, attr)) if attr else None)

        s.section("Terminal value & valuation", columns={"D": "Gordon Growth",
                                                         "E": "Exit Multiple"})
        sum_pv = lambda: f"=SUM({pv_range.format(row=s.rows['pv_fcf'])})"  # noqa: E731
        pair("sum_pv_fcf", "Sum of PV of free cash flows", sum_pv, sum_pv)
        pair("terminal_assumption", "Terminal growth rate / exit multiple",
             lambda: f"={inp('terminal_growth_rate')}", lambda: f"={inp('dcf_exit_multiple')}",
             fmts=(PCT, MULT))
        pair("terminal_metric", f"Final-year FCF / EBITDA (FY{last})",
             lambda: f"={s.r('unlevered_fcf', last)}",
             lambda: f"={s.r('ebitda', last, sheet='IS')}")
        pair("terminal_value", "Terminal value",
             lambda: f"={G('terminal_metric')}*(1+{G('terminal_assumption')})"
                     f"/({D('wacc')}-{G('terminal_assumption')})",
             lambda: f"={E('terminal_metric')}*{E('terminal_assumption')}",
             attr="terminal_value")
        tdf = lambda: f"=1/(1+{D('wacc')})^{s.r('period_index', last)}"  # noqa: E731
        pair("terminal_discount_factor", "Terminal discount factor (end of final year)",
             tdf, tdf, fmt=FACTOR)
        pair("pv_terminal_value", "PV of terminal value",
             lambda: f"={G('terminal_value')}*{G('terminal_discount_factor')}",
             lambda: f"={E('terminal_value')}*{E('terminal_discount_factor')}",
             attr="pv_terminal_value")
        pair("enterprise_value", "Enterprise value", style="subtotal",
             gordon=lambda: f"={G('sum_pv_fcf')}+{G('pv_terminal_value')}",
             exit_=lambda: f"={E('sum_pv_fcf')}+{E('pv_terminal_value')}",
             attr="enterprise_value")
        cash = lambda: f"={s.r('cash_and_investments', col=base, sheet='BS')}"  # noqa: E731
        pair("plus_cash", "Plus: cash & marketable securities", cash, cash)
        debt = lambda: f"=-{D('total_debt')}"  # noqa: E731
        pair("less_debt", "Less: total debt", debt, debt)
        tax = lambda: (f"=-IF({inp('deduct_repatriation_tax')}=1,"  # noqa: E731
                       f"{inp('foreign_cash')}*{inp('repatriation_tax_rate')},0)")
        pair("less_repatriation_tax", "Less: repatriation tax on offshore cash", tax, tax)
        pair("equity_value", "Implied equity value", style="subtotal",
             gordon=lambda: f"=SUM({s.span('enterprise_value', 'less_repatriation_tax', col='D')})",
             exit_=lambda: f"=SUM({s.span('enterprise_value', 'less_repatriation_tax', col='E')})",
             attr="equity_value")
        shares = lambda: f"={D('diluted_shares')}"  # noqa: E731
        pair("shares", "Diluted shares (m)", shares, shares, fmt=NUM1)
        pair("implied_share_price", "Implied share price", style="total", fmt=PRICE,
             gordon=lambda: f"={G('equity_value')}/{G('shares')}",
             exit_=lambda: f"={E('equity_value')}/{E('shares')}", attr="implied_share_price")
        pair("premium_to_current", "Premium / (discount) to current price", fmt=PCT,
             gordon=lambda: f"={G('implied_share_price')}/{D('share_price')}-1",
             exit_=lambda: f"={E('implied_share_price')}/{D('share_price')}-1",
             attr="premium_to_current")
        pair("pv_terminal_share_of_ev", "PV of terminal value, % of EV", fmt=PCT,
             gordon=lambda: f"={G('pv_terminal_value')}/{G('enterprise_value')}",
             exit_=lambda: f"={E('pv_terminal_value')}/{E('enterprise_value')}",
             attr="pv_terminal_share_of_ev")
        pair("implied_cross_check", "Implied exit multiple / implied perpetual growth",
             gordon=lambda: f"={G('terminal_value')}/{E('terminal_metric')}",
             exit_=lambda: f"=({E('terminal_value')}*{D('wacc')}-{G('terminal_metric')})"
                           f"/({E('terminal_value')}+{G('terminal_metric')})",
             fmts=(MULT, PCT), attr="implied_cross_check")

    # ================================================================= LBO
    def _declare_lbo(self) -> None:
        s = self.add(SheetSpec(self, "LBO", "Leveraged Buyout Analysis",
                               hist_years=[self.base_year], proj_years=self.lbo_years,
                               hist_label="Close", proj_label="Holding period"))
        lbo = self.model.lbo
        tx = self.a.transaction
        n = len(tx.tranches)
        base = self.col(self.base_year)
        first = self.lbo_years[0]
        last = self.lbo_years[-1]
        inp = lambda key: s.r(key, col="D", sheet="Inputs", absolute=True)  # noqa: E731
        inp_y = lambda key, y: s.r(key, y, sheet="Inputs")  # noqa: E731
        D = lambda key: s.r(key, col="D", absolute=True)  # noqa: E731
        sched = lbo.schedule
        exp = lambda key: (lambda y: float(sched.at[key, y])  # noqa: E731
                           if y in self.lbo_years else None)
        val = lambda mapping, k: (lambda _: mapping[k])  # noqa: E731

        s.section("Entry valuation")
        s.line("ltm_ebitda", f"LTM EBITDA (FY{self.base_year})",
               value=lambda: f"={s.r('ebitda', col=base, sheet='IS')}",
               expect=val(lbo.entry, "ltm_ebitda"))
        s.line("entry_multiple", "Entry EV / EBITDA", fmt=MULT,
               value=lambda: f"={inp('entry_multiple')}")
        s.line("entry_ev", "Entry enterprise value", style="subtotal",
               value=lambda: f"={D('ltm_ebitda')}*{D('entry_multiple')}",
               expect=val(lbo.entry, "entry_ev"))
        s.line("existing_debt", "Less: existing debt",
               value=lambda: f"={s.r('total_debt', col=base, sheet='BS')}")
        s.line("cash_and_investments", "Plus: cash & marketable securities",
               value=lambda: f"={s.r('cash_and_investments', col=base, sheet='BS')}")
        s.line("equity_purchase_price", "Equity purchase price", style="subtotal",
               value=lambda: f"={D('entry_ev')}-{D('existing_debt')}+{D('cash_and_investments')}",
               expect=val(lbo.entry, "equity_purchase_price"))
        s.line("diluted_shares", "Diluted shares (m)", fmt=NUM1,
               value=lambda: f"={inp('diluted_shares_outstanding')}")
        s.line("offer_price_per_share", "Implied offer price per share", fmt=PRICE,
               value=lambda: f"={D('equity_purchase_price')}/{D('diluted_shares')}",
               expect=val(lbo.entry, "offer_price_per_share"))
        s.line("offer_premium", "Premium to current share price", fmt=PCT,
               value=lambda: f"={D('offer_price_per_share')}/{inp('share_price')}-1",
               expect=val(lbo.entry, "offer_premium"))

        s.section("Cash available at close")
        s.line("foreign_cash", "Offshore cash & securities", value=lambda: f"={inp('foreign_cash')}")
        s.line("domestic_cash", "Domestic cash & securities",
               value=lambda: f"={D('cash_and_investments')}-{D('foreign_cash')}",
               expect=val(lbo.entry, "domestic_cash"))
        s.line("minimum_cash", "Minimum operating cash", value=lambda: f"={inp('minimum_cash')}")
        s.line("domestic_cash_used", "Domestic cash used to fund the deal",
               value=lambda: f"=MAX({D('domestic_cash')}-{D('minimum_cash')},0)",
               expect=val(lbo.entry, "domestic_cash_used"))
        s.line("opening_cash", "Cash retained on balance sheet",
               value=lambda: f"={D('domestic_cash')}-{D('domestic_cash_used')}",
               expect=val(lbo.entry, "opening_cash"))
        s.line("repatriated_cash", "Offshore cash repatriated",
               value=lambda: f"={D('foreign_cash')}*{inp('foreign_cash_usable_pct')}",
               expect=val(lbo.entry, "repatriated_cash"))
        s.line("repatriation_tax_paid", "Repatriation tax paid",
               value=lambda: f"={D('repatriated_cash')}*{inp('repatriation_tax_rate')}",
               expect=val(lbo.entry, "repatriation_tax_paid"))
        s.line("trapped_cash", "Offshore cash remaining trapped",
               value=lambda: f"={D('foreign_cash')}-{D('repatriated_cash')}",
               expect=val(lbo.entry, "trapped_cash"))

        s.section("Uses of funds")
        uses = list(lbo.uses)
        s.line("use_equity", uses[0], value=lambda: f"={D('equity_purchase_price')}")
        s.line("use_refinance", uses[1], value=lambda: f"={D('existing_debt')}")
        s.line("use_advisory", uses[2],
               value=lambda: f"={D('entry_ev')}*{inp('advisory_fee_pct')}",
               expect=val(lbo.uses, uses[2]))
        s.line("use_financing", uses[3],
               value=lambda: f"={D('total_new_debt')}*{inp('financing_fee_pct')}",
               expect=val(lbo.uses, uses[3]))
        s.line("total_uses", "Total uses", style="total",
               value=lambda: f"=SUM({s.span('use_equity', 'use_financing', col='D')})",
               expect=val(lbo.entry, "total_uses"))

        s.section("Sources of funds")
        s.line("total_new_debt", "Total new debt",
               value=lambda: f"={D('entry_ev')}*{inp('debt_pct_of_entry_ev')}",
               expect=val(lbo.entry, "total_new_debt"))
        for i, t in enumerate(tx.tranches):
            s.line(f"tranche{i}_principal", f"  {t.name}",
                   value=lambda i=i: f"={D('total_new_debt')}*{inp(f'tranche{i}_share')}",
                   expect=val(lbo.sources, t.name))
        s.line("source_cash", "Existing cash & securities used",
               value=lambda: f"={D('domestic_cash_used')}+{D('repatriated_cash')}"
                             f"-{D('repatriation_tax_paid')}",
               expect=val(lbo.sources, "Existing cash & securities used"))
        s.line("sponsor_equity", "Sponsor equity", style="subtotal",
               value=lambda: f"={D('total_uses')}-{D('total_new_debt')}-{D('source_cash')}",
               expect=val(lbo.sources, "Sponsor equity"))
        s.line("total_sources", "Total sources", style="total",
               value=lambda: f"=SUM({s.span('tranche0_principal', 'sponsor_equity', col='D')})",
               expect=val(lbo.entry, "total_sources"))
        s.line("sources_uses_check", "Check: sources - uses", style="check",
               value=lambda: f"={D('total_sources')}-{D('total_uses')}")
        s.line("debt_to_ebitda", "Total debt / LTM EBITDA", fmt=MULT,
               value=lambda: f"={D('total_new_debt')}/{D('ltm_ebitda')}",
               expect=val(lbo.entry, "debt_to_ebitda"))

        # ---- Annual cash flow (holding period columns)
        prev_or = lambda key, y, closing: (f"={closing}" if y == first  # noqa: E731
                                           else f"={s.r(key, y - 1)}")
        s.section("Cash flow available for debt service")
        s.line("ebitda", "EBITDA", proj=lambda y: f"={s.r('ebitda', y, sheet='IS')}",
               expect=exp("ebitda"))
        s.line("depreciation_and_amortization", "Less: D&A",
               proj=lambda y: f"={s.r('depreciation_and_amortization', y, sheet='IS')}")
        s.line("ebit", "EBIT", proj=lambda y: f"={s.r('ebitda', y)}-{s.r('depreciation_and_amortization', y)}",
               expect=exp("ebit"))
        s.line("interest_income", "Plus: interest income on cash",
               proj=lambda y: f"={inp_y('interest_income_yield', y)}*{s.r('beginning_cash', y)}",
               expect=exp("interest_income"))
        for i, t in enumerate(tx.tranches):
            s.line(f"tranche{i}_interest", f"Less: interest, {t.name}",
                   proj=lambda y, i=i: f"={inp(f'tranche{i}_rate')}*{s.r(f'tranche{i}_beginning', y)}",
                   expect=exp(f"tranche{i}_interest"))
        s.line("total_interest_expense", "Total interest expense", style="subtotal",
               proj=lambda y: "=" + "+".join(s.r(f"tranche{i}_interest", y) for i in range(n)),
               expect=exp("total_interest_expense"))
        s.line("pretax_income", "Pre-tax income",
               proj=lambda y: f"={s.r('ebit', y)}+{s.r('interest_income', y)}"
                              f"-{s.r('total_interest_expense', y)}",
               expect=exp("pretax_income"))
        s.line("tax_rate", "Tax rate", fmt=PCT, proj=lambda y: f"={inp_y('tax_rate', y)}")
        s.line("income_taxes", "Less: cash taxes (no NOL carry-forward)",
               proj=lambda y: f"=MAX({s.r('pretax_income', y)},0)*{s.r('tax_rate', y)}",
               expect=exp("income_taxes"))
        s.line("net_income", "Net income", style="subtotal",
               proj=lambda y: f"={s.r('pretax_income', y)}-{s.r('income_taxes', y)}",
               expect=exp("net_income"))
        s.line("capex", "Less: capital expenditures",
               proj=lambda y: f"=-{s.r('capital_expenditures', y, sheet='CFS')}",
               expect=exp("capex"))
        s.line("change_in_nwc", "Less: increase in net working capital",
               proj=lambda y: f"={s.r('change_in_nwc', y, sheet='DCF')}",
               expect=exp("change_in_nwc"))
        s.line("levered_fcf", "Levered free cash flow", style="total",
               proj=lambda y: f"={s.r('net_income', y)}+{s.r('depreciation_and_amortization', y)}"
                              f"-{s.r('capex', y)}-{s.r('change_in_nwc', y)}",
               expect=exp("levered_fcf"))

        s.section("Cash sweep")
        s.line("beginning_cash", "Cash, beginning of year",
               proj=lambda y: prev_or("ending_cash", y, D("opening_cash")),
               expect=exp("beginning_cash"))
        s.line("total_mandatory", "Less: mandatory amortization",
               proj=lambda y: "=" + "+".join(s.r(f"tranche{i}_mandatory", y) for i in range(n)),
               expect=exp("total_mandatory"))
        s.line("cash_available_for_sweep", "Cash available after minimum cash",
               proj=lambda y: f"={s.r('beginning_cash', y)}+{s.r('levered_fcf', y)}"
                              f"-{s.r('total_mandatory', y)}-{D('minimum_cash')}",
               expect=exp("cash_available_for_sweep"))
        s.line("sweep_pool", "Cash swept to optional prepayment",
               proj=lambda y: f"=MAX({s.r('cash_available_for_sweep', y)},0)*{inp('cash_sweep_pct')}",
               expect=exp("sweep_pool"))
        s.line("total_optional", "Optional prepayments",
               proj=lambda y: "=" + "+".join(s.r(f"tranche{i}_optional", y) for i in range(n)),
               expect=exp("total_optional"))
        s.line("ending_cash", "Cash, end of year", style="total",
               proj=lambda y: f"={s.r('beginning_cash', y)}+{s.r('levered_fcf', y)}"
                              f"-{s.r('total_mandatory', y)}-{s.r('total_optional', y)}",
               expect=exp("ending_cash"))

        for i, t in enumerate(tx.tranches):
            s.section(f"{t.name} ({t.interest_rate:.2%}, priority {i + 1})")
            s.line(f"tranche{i}_beginning", "Beginning balance",
                   proj=lambda y, i=i: prev_or(f"tranche{i}_ending", y, D(f"tranche{i}_principal")),
                   expect=exp(f"tranche{i}_beginning"))
            s.line(f"tranche{i}_mandatory", "Mandatory amortization",
                   proj=lambda y, i=i: f"=MIN({s.r(f'tranche{i}_beginning', y)},"
                                       f"{D(f'tranche{i}_principal')}*{inp(f'tranche{i}_amort')})",
                   expect=exp(f"tranche{i}_mandatory"))

            def optional(y: int, i: int = i) -> str:
                prior = "".join(f"-{s.r(f'tranche{j}_optional', y)}" for j in range(i))
                return (f"=IF({inp(f'tranche{i}_prepayable')}=1,MIN({s.r('sweep_pool', y)}{prior},"
                        f"{s.r(f'tranche{i}_beginning', y)}-{s.r(f'tranche{i}_mandatory', y)}),0)")

            s.line(f"tranche{i}_optional", "Optional prepayment (sweep)", proj=optional,
                   expect=exp(f"tranche{i}_optional"))
            s.line(f"tranche{i}_ending", "Ending balance", style="subtotal",
                   proj=lambda y, i=i: f"={s.r(f'tranche{i}_beginning', y)}"
                                       f"-{s.r(f'tranche{i}_mandatory', y)}"
                                       f"-{s.r(f'tranche{i}_optional', y)}",
                   expect=exp(f"tranche{i}_ending"))

        s.section("Credit statistics & covenants")
        s.line("total_debt", "Total debt", style="subtotal",
               proj=lambda y: "=" + "+".join(s.r(f"tranche{i}_ending", y) for i in range(n)),
               expect=exp("total_debt"))
        s.line("net_debt", "Net debt",
               proj=lambda y: f"={s.r('total_debt', y)}-{s.r('ending_cash', y)}",
               expect=exp("net_debt"))
        s.line("leverage", "Total debt / EBITDA", fmt=MULT,
               proj=lambda y: f"={s.r('total_debt', y)}/{s.r('ebitda', y)}",
               expect=exp("leverage"))
        s.line("interest_coverage", "EBITDA / interest expense", fmt=MULT,
               proj=lambda y: f"=IF({s.r('total_interest_expense', y)}=0,\"n.m.\","
                              f"{s.r('ebitda', y)}/{s.r('total_interest_expense', y)})",
               expect=exp("interest_coverage"))
        s.line("leverage_covenant", "Covenant: max leverage", fmt=MULT,
               proj=lambda y: f"={inp('max_leverage_covenant')}")
        s.line("leverage_cushion", "EBITDA cushion to leverage covenant", fmt=PCT,
               proj=lambda y: f"=1-({s.r('total_debt', y)}/{s.r('leverage_covenant', y)})"
                              f"/{s.r('ebitda', y)}",
               expect=exp("leverage_cushion"))
        s.line("coverage_covenant", "Covenant: min interest coverage", fmt=MULT,
               proj=lambda y: f"={inp('min_coverage_covenant')}")
        s.line("coverage_cushion", "EBITDA cushion to coverage covenant", fmt=PCT,
               proj=lambda y: f"=1-{s.r('coverage_covenant', y)}*{s.r('total_interest_expense', y)}"
                              f"/{s.r('ebitda', y)}",
               expect=exp("coverage_cushion"))
        s.line("covenant_status", "Covenant compliance", fmt=TEXT,
               proj=lambda y: f"=IF(AND({s.r('leverage', y)}<={s.r('leverage_covenant', y)},"
                              f"{s.r('coverage_cushion', y)}>=0),\"Pass\",\"BREACH\")")

        s.section("Exit & sponsor returns")
        s.line("exit_ebitda", f"Exit-year EBITDA (FY{last})",
               value=lambda: f"={s.r('ebitda', last)}", expect=val(lbo.exit, "exit_ebitda"))
        s.line("exit_multiple", "Exit EV / EBITDA", fmt=MULT,
               value=lambda: f"={inp('exit_multiple')}")
        s.line("exit_ev", "Exit enterprise value", style="subtotal",
               value=lambda: f"={D('exit_ebitda')}*{D('exit_multiple')}",
               expect=val(lbo.exit, "exit_ev"))
        s.line("exit_debt", "Less: debt outstanding", value=lambda: f"={s.r('total_debt', last)}",
               expect=val(lbo.exit, "exit_debt"))
        s.line("exit_cash", "Plus: cash", value=lambda: f"={s.r('ending_cash', last)}",
               expect=val(lbo.exit, "exit_cash"))
        s.line("trapped_cash_after_tax", "Plus: trapped offshore cash, after repatriation tax",
               value=lambda: f"={D('trapped_cash')}*(1-{inp('repatriation_tax_rate')})",
               expect=val(lbo.exit, "trapped_cash_after_tax"))
        s.line("exit_equity", "Exit equity value", style="total",
               value=lambda: f"={D('exit_ev')}-{D('exit_debt')}+{D('exit_cash')}"
                             f"+{D('trapped_cash_after_tax')}",
               expect=val(lbo.exit, "exit_equity"))
        s.line("sponsor_cash_flow", "Sponsor cash flows", style="subtotal",
               cells={base: (lambda: f"=-{D('sponsor_equity')}", NUM)},
               proj=lambda y: f"=IF({self.col(y)}${HEADER_ROW}={self.col(last)}${HEADER_ROW},"
                              f"{D('exit_equity')},0)",
               expect=lambda c: (lbo.sponsor_cash_flows[0] if c == base else
                                 lbo.sponsor_cash_flows[self.lbo_years.index(c) + 1]
                                 if c in self.lbo_years else None))
        s.line("moic", "MOIC", fmt=MULT, style="total",
               value=lambda: f"={D('exit_equity')}/{D('sponsor_equity')}",
               expect=lambda _: lbo.moic)
        s.line("irr", "IRR", fmt=PCT, style="total",
               value=lambda: f"=IRR({base}{s.rows['sponsor_cash_flow']}:"
                             f"{self.col(last)}{s.rows['sponsor_cash_flow']})",
               expect=lambda _: lbo.irr)

    # ================================================================= rendering
    def _render(self, ws: Worksheet, spec: SheetSpec) -> None:
        """Write one declared sheet: headers, labels, cells, styles, expectations."""
        ws.sheet_view.showGridLines = True
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["C"].width = 9
        for i in range(FIRST_YEAR_COL, FIRST_YEAR_COL + len(self.all_years) + 1):
            ws.column_dimensions[col_letter(i)].width = 13
        ws["B1"] = f"{self.schema.metadata.name}: {spec.title}"
        ws["B1"].font = Font(name=FONT_NAME, size=14, bold=True, color=NAVY)
        ws["B2"] = ("USD in millions, except per-share data. Blue = input / reported, "
                    "black = formula, green = link to another sheet.")
        ws["B2"].font = Font(name=FONT_NAME, size=9, italic=True, color="595959")

        if spec.hist_years or spec.proj_years:
            self._render_year_header(ws, spec)
        ws.freeze_panes = ws.cell(row=FIRST_BODY_ROW, column=FIRST_YEAR_COL)

        last_col = FIRST_YEAR_COL + len(self.all_years) - 1
        for line in spec.lines:
            if line.kind == "section":
                for c in range(2, last_col + 1):
                    cell = ws.cell(row=line.row, column=c)
                    cell.fill = PatternFill("solid", fgColor=SECTION_FILL)
                    cell.font = Font(name=FONT_NAME, bold=True, color=NAVY)
                ws.cell(row=line.row, column=2, value=line.label)
                for col, text in spec.column_headers.get(line.row, {}).items():
                    cell = ws[f"{col}{line.row}"]
                    cell.value = text
                    cell.alignment = Alignment(horizontal="right")
                continue
            if line.kind == "blank":
                continue
            label = ws.cell(row=line.row, column=2, value=line.label)
            label.font = Font(name=FONT_NAME, bold=line.style in ("subtotal", "total"),
                              italic=line.style in ("memo", "check"),
                              color="C00000" if line.style == "check" else "000000")
            if line.unit:
                ws.cell(row=line.row, column=3, value=line.unit).font = Font(
                    name=FONT_NAME, size=9, color="595959")

            targets: list[tuple[str, int | str, CellValue, str]] = []
            for y in spec.hist_years:
                if line.hist is not None:
                    targets.append((self.col(y), y, line.hist(y), line.fmt))
            for y in spec.proj_years:
                if line.proj is not None:
                    targets.append((self.col(y), y, line.proj(y), line.fmt))
            for col, (fn, fmt) in line.cells.items():
                targets.append((col, col, fn(), fmt))

            for col, colkey, value, fmt in targets:
                if value is None:
                    continue
                cell = ws[f"{col}{line.row}"]
                cell.value = value
                cell.number_format = fmt
                self._style_cell(cell, value, line.style, spec.inputs_sheet)
                if line.style == "check":
                    suffix = f" FY{colkey}" if isinstance(colkey, int) else ""
                    self.checks.append((f"{spec.name}: {line.label}{suffix}", spec.name,
                                        f"{col}{line.row}"))
                if line.expect is not None and isinstance(value, str) and value.startswith("="):
                    expected = line.expect(colkey)
                    if expected is not None:
                        self.expected[f"{spec.name}!{col}{line.row}"] = float(expected)

        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def _render_year_header(self, ws: Worksheet, spec: SheetSpec) -> None:
        """Banner row (Historical / Projected) and numeric fiscal-year header row."""
        def banner(years: list[int], text: str, color: str) -> None:
            if not years:
                return
            c1, c2 = self.col(years[0]), self.col(years[-1])
            ws.merge_cells(f"{c1}{HEADER_ROW - 1}:{c2}{HEADER_ROW - 1}")
            cell = ws[f"{c1}{HEADER_ROW - 1}"]
            cell.value = text
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
            for y in years:
                ws[f"{self.col(y)}{HEADER_ROW - 1}"].fill = PatternFill("solid", fgColor=color)

        banner(spec.hist_years, spec.hist_label, NAVY)
        banner(spec.proj_years, spec.proj_label, BLUE_HDR)
        label = ws.cell(row=HEADER_ROW, column=2, value="Fiscal year ending September")
        label.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        for c in range(2, FIRST_YEAR_COL + len(self.all_years)):
            ws.cell(row=HEADER_ROW, column=c).fill = PatternFill("solid", fgColor=NAVY)
        for years, suffix in ((spec.hist_years, "A"), (spec.proj_years, "E")):
            for y in years:
                cell = ws[f"{self.col(y)}{HEADER_ROW}"]
                cell.value = y
                cell.number_format = f'"FY"0"{suffix}"'
                cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="right")

    @staticmethod
    def _style_cell(cell: Any, value: CellValue, style: str, inputs_sheet: bool) -> None:
        """Apply the font colour convention and subtotal/total borders."""
        is_formula = isinstance(value, str) and value.startswith("=")
        if is_formula:
            color = "008000" if "!" in value else "000000"
        elif isinstance(value, (int, float)):
            color = "0000FF"
            if inputs_sheet:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
        else:
            color = "000000"
        if style == "check":
            color = "C00000"
        cell.font = Font(name=FONT_NAME, color=color, bold=style in ("subtotal", "total"),
                         italic=style == "memo")
        if style == "subtotal":
            cell.border = Border(top=THIN)
        elif style == "total":
            cell.border = Border(top=THIN, bottom=DOUBLE)
        if isinstance(value, str) and not is_formula:
            cell.alignment = Alignment(horizontal="left")

    def _render_checks(self, ws: Worksheet) -> None:
        """One row per identity check, a pass/fail flag, and a master check cell."""
        ws.sheet_view.showGridLines = True
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 60
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 10
        ws["B1"] = f"{self.schema.metadata.name}: Model Integrity Checks"
        ws["B1"].font = Font(name=FONT_NAME, size=14, bold=True, color=NAVY)
        ws["B2"] = "Every check must be zero (within $0.5M). Flag 1 = error."
        ws["B2"].font = Font(name=FONT_NAME, size=9, italic=True, color="595959")
        first, last = 6, 6 + len(self.checks) - 1
        ws["B4"] = "MASTER CHECK"
        ws["B4"].font = Font(name=FONT_NAME, bold=True)
        ws["C4"] = f'=IF(SUM(D{first}:D{last})=0,"OK","ERROR")'
        ws["C4"].font = Font(name=FONT_NAME, bold=True, color="008000")
        for c, text in (("B", "Check"), ("C", "Difference"), ("D", "Flag")):
            cell = ws[f"{c}5"]
            cell.value = text
            cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
        for i, (label, sheet, addr) in enumerate(self.checks):
            r = first + i
            ws[f"B{r}"] = label
            ws[f"C{r}"] = f"='{sheet}'!{addr}"
            ws[f"C{r}"].number_format = NUM1
            ws[f"C{r}"].font = Font(name=FONT_NAME, color="008000")
            ws[f"D{r}"] = f"=IF(ABS(C{r})<0.5,0,1)"
            ws[f"D{r}"].number_format = INT
            self.expected[f"Checks!C{r}"] = 0.0
        self.expected["Checks!C4"] = "OK"  # type: ignore[assignment]
        self.master_check = "Checks!C4"

    def _render_cover(self, ws: Worksheet) -> None:
        """Title page: transaction summary, key outputs (linked), contents, conventions."""
        md, d, lbo = self.schema.metadata, self.model.dcf, self.model.lbo
        ws.sheet_view.showGridLines = True
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 46
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 18
        ws["B2"] = f"{md.name}: Valuation Model"
        ws["B2"].font = Font(name=FONT_NAME, size=20, bold=True, color=NAVY)
        ws["B3"] = ("3-statement projection, DCF and hypothetical take-private LBO stress test")
        ws["B3"].font = Font(name=FONT_NAME, size=11, italic=True, color="595959")
        expected_by_label = {
            "DCF implied share price: Gordon Growth": d.gordon.implied_share_price,
            "DCF implied share price: Exit Multiple": d.exit.implied_share_price,
            "Sponsor MOIC": lbo.moic,
            "Sponsor IRR": lbo.irr,
        }
        rows: list[tuple[str, Any, str]] = [
            ("Source filing", f"{md.filing_type}, fiscal year ended {md.fiscal_year_end:%d %B %Y}", TEXT),
            ("SEC CIK", md.cik or "n/a", TEXT),
            ("Model generated", datetime.now().strftime("%d %B %Y"), TEXT),
            ("", None, TEXT),
            ("KEY OUTPUTS (linked)", None, TEXT),
            ("Current share price", f"={self.ref('Inputs', 'share_price', col='D')}", PRICE),
            ("DCF implied share price: Gordon Growth",
             f"={self.ref('DCF', 'implied_share_price', col='D')}", PRICE),
            ("DCF implied share price: Exit Multiple",
             f"={self.ref('DCF', 'implied_share_price', col='E')}", PRICE),
            ("WACC", f"={self.ref('DCF', 'wacc', col='D')}", PCT2),
            ("LBO entry EV / EBITDA", f"={self.ref('LBO', 'entry_multiple', col='D')}", MULT),
            ("LBO implied offer price", f"={self.ref('LBO', 'offer_price_per_share', col='D')}",
             PRICE),
            ("Sponsor equity cheque", f"={self.ref('LBO', 'sponsor_equity', col='D')}", NUM),
            ("Sponsor MOIC", f"={self.ref('LBO', 'moic', col='D')}", MULT),
            ("Sponsor IRR", f"={self.ref('LBO', 'irr', col='D')}", PCT),
            ("Model integrity (all checks)", "=Checks!C4", TEXT),
        ]
        for i, (label, value, fmt) in enumerate(rows, start=5):
            ws[f"B{i}"] = label
            ws[f"B{i}"].font = Font(name=FONT_NAME, bold=label.isupper() and bool(label))
            if value is not None:
                cell = ws[f"C{i}"]
                cell.value = value
                cell.number_format = fmt
                cell.font = Font(name=FONT_NAME, bold=True,
                                 color="008000" if str(value).startswith("=") else "000000")
                cell.alignment = Alignment(horizontal="right")
                if label in expected_by_label:
                    self.expected[f"Cover!C{i}"] = expected_by_label[label]

        start = 5 + len(rows) + 1
        notes = [
            "CONTENTS",
            "Inputs: market, operating and transaction assumptions (all hard-coded inputs live here)",
            "IS / BS / CFS: linked three-statement model, historical and projected",
            "Debt: existing debt tranches from the 10-K, reconciled to the balance sheet",
            "DCF: WACC, unlevered FCF, Gordon Growth and Exit Multiple valuation",
            "LBO: sources & uses, cash sweep waterfall, covenants, sponsor returns",
            "Checks: every accounting identity, with a master check",
            "",
            "CONVENTIONS",
            "Interest is calculated on beginning-of-year balances (no circular references).",
            "Funded debt is refinanced (held flat) in the standalone projection.",
            "Offshore cash is valued net of the repatriation tax implied by the 10-K.",
            f"Apple is far beyond any realised LBO size; the LBO is a hypothetical stress test.",
        ]
        for i, text in enumerate(notes):
            cell = ws[f"B{start + i}"]
            cell.value = text
            cell.font = Font(name=FONT_NAME, bold=text.isupper() and bool(text), size=10)


@dataclass(frozen=True)
class WorkbookReport:
    """Result of :func:`write_workbook`.

    Attributes:
        path: Where the workbook was written.
        expected: Python-engine value for each formula cell, keyed ``"Sheet!A1"``.
        formula_cells: Number of formula cells written.
        input_cells: Number of hard-coded numeric cells written.
        checks: Number of identity checks on the Checks sheet.
    """

    path: Path
    expected: dict[str, Any]
    formula_cells: int
    input_cells: int
    checks: int


def write_workbook(model: ValuationModel, path: str | Path) -> WorkbookReport:
    """Render the full model to an ``.xlsx`` workbook.

    Args:
        model: Output of :func:`src.modeling_engine.build_model`.
        path: Destination file.

    Returns:
        A report including the expected value of every formula cell, used by the
        test-suite to reconcile Excel against the Python engine.
    """
    builder = WorkbookBuilder(model)
    wb = builder.build()
    # openpyxl stores formulas without cached results: force Excel to calculate on open.
    wb.calculation.fullCalcOnLoad = True
    path = Path(path)
    wb.save(path)

    formulas = inputs = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                elif isinstance(cell.value, (int, float)) and cell.row > HEADER_ROW:
                    inputs += 1
    return WorkbookReport(path=path, expected=builder.expected, formula_cells=formulas,
                          input_cells=inputs, checks=len(builder.checks))
