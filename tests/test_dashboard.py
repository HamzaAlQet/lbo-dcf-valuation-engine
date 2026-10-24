"""Tests for the Streamlit dashboard: pure helpers plus headless app runs (AppTest)."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from pydantic import ValidationError

from src import memo_generator
from src.analyzer import AnalysisResult
from src.dashboard import (
    Controls,
    apply_controls,
    covenant_frame,
    heatmap,
    scenario_frame,
    sources_uses_frame,
    workbook_bytes,
)
from src.memo_generator import Memo

pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest  # noqa: E402

APP = str(Path(__file__).resolve().parents[1] / "streamlit_app.py")


# --------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------
def test_default_controls_reproduce_base_assumptions(analysis: AnalysisResult) -> None:
    base = analysis.base.model.assumptions
    assert apply_controls(base, Controls.from_assumptions(base)) == base


def test_shifts_apply_to_every_year(analysis: AnalysisResult) -> None:
    base = analysis.base.model.assumptions
    controls = Controls(**{**Controls.from_assumptions(base).as_dict(),
                           "revenue_growth_shift": 0.01, "entry_multiple": 7.0})
    adjusted = apply_controls(base, controls)
    assert adjusted.operating.revenue_growth == pytest.approx(
        [g + 0.01 for g in base.operating.revenue_growth])
    assert adjusted.transaction.entry_ev_ebitda_multiple == 7.0


def test_invalid_combination_is_rejected(analysis: AnalysisResult) -> None:
    base = analysis.base.model.assumptions
    bad = Controls(**{**Controls.from_assumptions(base).as_dict(), "gross_margin_shift": -0.35})
    with pytest.raises(ValidationError):
        apply_controls(base, bad)


def test_tables_and_charts_build(analysis: AnalysisResult) -> None:
    uses, sources = sources_uses_frame(analysis.base.model.lbo)
    assert uses["$B"].iloc[-1] == pytest.approx(sources["$B"].iloc[-1])  # S = U
    assert list(scenario_frame(analysis).columns) == ["Base", "Bull", "Bear"]
    assert len(covenant_frame(analysis.base.model.lbo)) == 5
    spec = heatmap(analysis.sensitivities.lbo_irr).to_dict()
    assert spec["layer"]  # cells, base-case outline, labels


def test_workbook_bytes_is_a_valid_workbook(analysis: AnalysisResult) -> None:
    wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes(analysis)))
    assert "Sensitivity" in wb.sheetnames and "Scenarios" in wb.sheetnames


def test_strict_mode_raises_instead_of_falling_back(analysis: AnalysisResult) -> None:
    class NoCredentials:
        class beta:  # noqa: N801
            class messages:  # noqa: N801
                @staticmethod
                def parse(**_: Any) -> Any:
                    raise TypeError("Could not resolve authentication method.")

    with pytest.raises(memo_generator.MemoGenerationError, match="no API key"):
        memo_generator.generate_memo(analysis, client=NoCredentials(),
                                     allow_template_fallback=False)


# --------------------------------------------------------------------------------------
# Headless app runs
# --------------------------------------------------------------------------------------
@pytest.fixture
def app(analysis: AnalysisResult) -> AppTest:  # `analysis` warms the extraction cache path
    at = AppTest.from_file(APP, default_timeout=300)
    at.run()
    assert not at.exception, at.exception
    return at


def test_app_shows_base_case_and_sample_memo(app: AppTest) -> None:
    assert "Recommendation: PASS" in app.error[0].value
    labels = [m.label for m in app.metric]
    assert "Sponsor IRR" in labels and "DCF: Gordon Growth" in labels
    irr = next(m for m in app.metric if m.label == "Sponsor IRR")
    assert irr.value == "7.5%"
    assert any("Sample memo (base case)" in c.value for c in app.caption)
    assert not app.info  # no "custom assumptions" banner by default


def test_flexing_assumptions_recomputes(app: AppTest) -> None:
    app.slider(key="entry_multiple").set_value(6.0).run()
    assert not app.exception, app.exception
    irr = next(m for m in app.metric if m.label == "Sponsor IRR")
    assert irr.value == "14.1%"  # matches the 6.0x / 50% cell of the IRR grid
    assert app.info  # custom-assumptions banner
    assert any("Template memo for these assumptions" in c.value for c in app.caption)


def test_ticker_mode_values_another_company(monkeypatch: pytest.MonkeyPatch) -> None:
    import json as _json

    from src import xbrl_mapping
    from src.xbrl import company_facts_to_filing

    fixture = Path(__file__).parent / "fixtures" / "companyfacts_HD.json"
    hd = xbrl_mapping.build_schema(company_facts_to_filing(
        _json.loads(fixture.read_text(encoding="utf-8")), ticker="HD"))
    calls: list[tuple[str, int | None]] = []

    def fake_lookup(ticker: str, fiscal_year: int | None = None) -> Any:
        calls.append((ticker, fiscal_year))
        return hd

    monkeypatch.setattr(xbrl_mapping, "schema_from_ticker", fake_lookup)
    at = AppTest.from_file(APP, default_timeout=300)
    at.run()
    at.radio(key="mode").set_value("Look up a ticker").run()
    assert not at.exception, at.exception
    at.text_input(key="ticker_input").input("hd")
    at.number_input(key="price_ticker").set_value(380.0)
    at.button[0].click().run()  # the form's submit button
    assert not at.exception, at.exception
    assert calls == [("HD", None)]
    assert "HOME DEPOT" in at.title[0].value.upper()
    labels = {m.label for m in at.metric}
    assert {"Sponsor IRR", "LBO offer price"} <= labels
    offer = next(m for m in at.metric if m.label == "LBO offer price")
    assert "+30.0% vs $380.00" in offer.delta


def test_upload_mode_values_an_inline_xbrl_10k() -> None:
    at = AppTest.from_file(APP, default_timeout=300)
    at.run()
    at.radio(key="mode").set_value("Upload a 10-K (.htm)").run()
    assert not at.exception, at.exception
    assert any("Enter a ticker or upload" in i.value for i in at.info)  # waits for input


def test_claude_memo_uses_visitor_key_not_server_key(app: AppTest,
                                                     monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "server-key-must-not-be-used")
    seen: dict[str, Any] = {}
    real = memo_generator.generate_memo

    def fake(analysis: AnalysisResult, **kwargs: Any) -> Memo:
        if kwargs.get("offline"):
            return real(analysis, **kwargs)
        seen["api_key"] = kwargs["client"].api_key
        seen["strict"] = kwargs["allow_template_fallback"] is False
        return Memo(markdown="# Memo drafted by the stub", source="claude-opus-5-5")

    monkeypatch.setattr(memo_generator, "generate_memo", fake)
    assert app.button(key="generate_memo").disabled  # no key entered yet
    app.text_input(key="api_key").input("sk-ant-visitor").run()
    app.button(key="generate_memo").click().run()
    assert not app.exception, app.exception
    assert seen == {"api_key": "sk-ant-visitor", "strict": True}
    assert any("drafted by Claude" in c.value for c in app.caption)
    assert any("Memo drafted by the stub" in m.value for m in app.markdown)
