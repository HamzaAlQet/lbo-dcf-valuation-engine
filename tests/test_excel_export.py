"""Excel workbook tests.

The core test evaluates every formula in the generated workbook with the
``formulas`` library (a pure-Python Excel calculation engine) and reconciles each
result against the Python engine. This proves the workbook is a faithful, live
version of the model rather than pasted values.
"""

from __future__ import annotations

import math
import re
import warnings
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from src.excel_export import HEADER_ROW, WorkbookReport, write_workbook
from src.modeling_engine import build_model
from src.schemas import FinancialStatementSchema, ModelAssumptions

formulas = pytest.importorskip("formulas")

CONFIG = Path(__file__).resolve().parents[1] / "config" / "assumptions.json"
CALC_SHEETS_PROJECTED = ("IS", "BS", "CFS")
CALC_SHEETS_ALL = ("DCF", "LBO", "Checks")


def evaluate(path: Path) -> Callable[[str], Any]:
    """Calculate a workbook; return a lookup ``"Sheet!A1" -> value``."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    values: dict[str, Any] = {}
    for key, ranges in solution.items():
        m = re.match(r"^'\[.+?\](.+?)'!([A-Z]+\d+)$", key)
        if m:
            values[f"{m.group(1)}!{m.group(2)}".upper()] = ranges.value[0][0]
    return lambda addr: values.get(addr.upper())


@pytest.fixture(scope="module")
def workbook(apple: FinancialStatementSchema,
             tmp_path_factory: pytest.TempPathFactory) -> tuple[WorkbookReport, Callable[[str], Any]]:
    model = build_model(apple, ModelAssumptions.from_json(CONFIG))
    report = write_workbook(model, tmp_path_factory.mktemp("xlsx") / "model.xlsx")
    return report, evaluate(report.path)


def test_every_formula_reconciles_to_python(
        workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    report, value = workbook
    assert len(report.expected) > 500
    mismatches = []
    for addr, expected in report.expected.items():
        got = value(addr)
        if isinstance(expected, str):
            ok = got == expected
        else:
            ok = isinstance(got, (int, float)) and math.isclose(
                float(got), expected, rel_tol=1e-9, abs_tol=1e-6)
        if not ok:
            mismatches.append(f"{addr}: excel={got!r} python={expected!r}")
    assert not mismatches, "\n".join(mismatches[:20])


def test_master_check_passes(workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    report, value = workbook
    assert report.checks >= 20
    assert value("Checks!C4") == "OK"


def test_broken_formula_trips_master_check(
        workbook: tuple[WorkbookReport, Callable[[str], Any]], tmp_path: Path) -> None:
    """Negative control: add $10M (well beyond the $1M rounding tolerance) to one
    projected total assets cell."""
    report, _ = workbook
    wb = openpyxl.load_workbook(report.path)
    ws = wb["BS"]
    row = next(c.row for c in ws["B"] if c.value == "Total assets")
    ws[f"H{row}"].value += "+10"
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    assert evaluate(broken)("Checks!C4") == "ERROR"


def test_no_hardcoded_numbers_in_calculations(
        workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    """Projection columns of the statements, and every DCF / LBO / Checks cell,
    must be formulas. The only hard-coded numbers are the fiscal-year headers."""
    report, _ = workbook
    wb = openpyxl.load_workbook(report.path)
    offenders = []
    for name in CALC_SHEETS_PROJECTED + CALC_SHEETS_ALL:
        ws = wb[name]
        for row in ws.iter_rows(min_row=HEADER_ROW + 1):
            for cell in row:
                projected = cell.column >= 7  # columns G+ hold projection years
                if (isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                        and (name in CALC_SHEETS_ALL or projected)):
                    offenders.append(f"{name}!{cell.coordinate}={cell.value}")
    assert not offenders, offenders[:10]


def test_formulas_use_uppercase_functions(
        workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    report, _ = workbook
    wb = openpyxl.load_workbook(report.path)
    functions = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    functions.update(re.findall(r"([A-Za-z]+)\(", cell.value))
    assert functions >= {"SUM", "MIN", "MAX", "IF", "IRR", "SUMPRODUCT"}
    assert all(f.isupper() for f in functions), functions


@pytest.fixture(scope="module")
def phase3_workbook(analysis: Any, tmp_path_factory: pytest.TempPathFactory
                    ) -> tuple[WorkbookReport, Callable[[str], Any]]:
    report = write_workbook(analysis.base.model,
                            tmp_path_factory.mktemp("xlsx3") / "model.xlsx", analysis=analysis)
    return report, evaluate(report.path)


def test_sensitivity_formulas_reconcile_to_python(
        phase3_workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    report, value = phase3_workbook
    sens = {a: e for a, e in report.expected.items() if a.startswith("Sensitivity!")}
    assert len(sens) >= 60  # two 5x5 DCF grids plus their axes
    bad = [a for a, e in sens.items()
           if not math.isclose(float(value(a)), e, rel_tol=1e-9, abs_tol=1e-9)]
    assert not bad, bad[:10]
    assert value("Checks!C4") == "OK"


def test_phase3_sheets_present(
        phase3_workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    report, _ = phase3_workbook
    wb = openpyxl.load_workbook(report.path)
    assert wb.sheetnames[-3:] == ["Sensitivity", "Scenarios", "Checks"]
    ws = wb["Sensitivity"]
    assert len(ws.conditional_formatting) == 4  # one heat map per grid
    live_dcf_cells = [c for row in ws.iter_rows() for c in row
                      if isinstance(c.value, str) and "SUMPRODUCT" in c.value]
    assert len(live_dcf_cells) == 50  # every DCF grid cell is a formula


def test_colour_convention_and_layout(
        workbook: tuple[WorkbookReport, Callable[[str], Any]]) -> None:
    report, _ = workbook
    wb = openpyxl.load_workbook(report.path)
    assert wb.sheetnames == ["Cover", "Inputs", "IS", "BS", "CFS", "Debt", "DCF", "LBO", "Checks"]
    inputs = wb["Inputs"]
    row = next(c.row for c in inputs["B"] if c.value == "Risk-free rate (10Y Treasury)")
    assert inputs[f"D{row}"].font.color.rgb.endswith("0000FF")      # input: blue
    is_ = wb["IS"]
    row = next(c.row for c in is_["B"] if c.value == "Net sales")
    assert is_[f"G{row}"].font.color.rgb.endswith("008000")          # cross-sheet link: green
    row = next(c.row for c in is_["B"] if c.value == "Gross profit")
    assert is_[f"G{row}"].font.color.rgb.endswith("000000")          # same-sheet formula: black
    for ws in wb.worksheets:
        assert ws.sheet_view.showGridLines
